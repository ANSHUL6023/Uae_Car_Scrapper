#!/usr/bin/env python3
"""
automate.py — Excel-driven scrape → WordPress upload pipeline (UAE / zigwheels.ae)

Reads a spreadsheet of brands/models, scrapes each car one by one from
zigwheels.ae with scrapper.py, marks scrape status, uploads to WordPress
(uae.fyndmycar.com) with wp_uploader.py, and marks upload status.

Excel columns (header row required):
  brand, model, url, scrape_status, upload_status, acf_json, error, scraped_at, uploaded_at

  - brand / model : required (used to build URL if url is empty)
  - url           : optional override, e.g. https://www.zigwheels.ae/new-cars/hyundai/creta
  - scrape_status : pending | done | error | skipped
  - upload_status : pending | done | error | skipped
  - acf_json      : filled automatically after a successful scrape
  - error         : last error message (cleared on retry)

Usage
-----
  # Create a starter spreadsheet
  python automate.py --create-template cars.xlsx

  # Run full pipeline (scrape then upload pending rows)
  python automate.py cars.xlsx --remove-logo --rephrase-content

  # Credentials from .env (WP_URL, WP_USER, WP_PASSWORD) or CLI flags
  python automate.py cars.xlsx \\
      --wp-url https://uae.fyndmycar.com/ \\
      --user FyndMyCar_admin \\
      --password "xxxx xxxx xxxx xxxx"

  # Scrape only / upload only
  python automate.py cars.xlsx --scrape-only --remove-logo --rephrase-content
  python automate.py cars.xlsx --upload-only
"""

from __future__ import annotations

import argparse
import logging
import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv
import os

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    print("ERROR: openpyxl is required. Run: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

load_dotenv()

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-7s  %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger(__name__)

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_OUTDIR = SCRIPT_DIR / "pipeline_output"
ZIGWHEELS_BASE = "https://www.zigwheels.ae"

COLUMNS = [
    "brand",
    "model",
    "url",
    "scrape_status",
    "upload_status",
    "acf_json",
    "error",
    "scraped_at",
    "uploaded_at",
]

SAMPLE_ROWS = [
    ("Hyundai", "Creta", "", "pending", "pending", "", "", "", ""),
    ("Toyota", "Fortuner", "", "pending", "pending", "", "", "", ""),
    ("Nissan", "Patrol", "", "pending", "pending", "", "", "", ""),
]


def slugify(text: str) -> str:
    text = text.lower().strip()
    text = re.sub(r"[^a-z0-9\s-]", "", text)
    return re.sub(r"\s+", "-", text)


def build_url(brand: str, model: str) -> str:
    brand_slug = slugify(brand)
    model_slug = slugify(model)
    return f"{ZIGWHEELS_BASE}/new-cars/{brand_slug}/{model_slug}"


def output_stem(brand: str, model: str) -> str:
    return f"{slugify(brand)}-{slugify(model)}"


def now_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def create_template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "cars"
    ws.append(COLUMNS)
    for row in SAMPLE_ROWS:
        ws.append(list(row))
    wb.save(path)
    log.info("Template created → %s", path.resolve())


def _cell(row: tuple, col_index: dict, name: str) -> str:
    idx = col_index.get(name)
    if idx is None:
        return ""
    val = row[idx]
    return "" if val is None else str(val).strip()


def load_jobs(path: Path) -> tuple[list[dict], object, dict]:
    wb = load_workbook(path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Spreadsheet is empty: {path}")

    headers = [str(h).strip().lower() if h else "" for h in rows[0]]
    col_index = {name: i for i, name in enumerate(headers) if name}

    missing = [c for c in ("brand", "model") if c not in col_index]
    if missing:
        raise ValueError(f"Missing required column(s): {', '.join(missing)}")

    jobs = []
    for excel_row, row in enumerate(rows[1:], start=2):
        brand = _cell(row, col_index, "brand")
        model = _cell(row, col_index, "model")
        if not brand and not model:
            continue
        if not brand or not model:
            log.warning("Row %d skipped — brand and model are both required", excel_row)
            continue

        jobs.append(
            {
                "excel_row": excel_row,
                "brand": brand,
                "model": model,
                "url": _cell(row, col_index, "url"),
                "scrape_status": _cell(row, col_index, "scrape_status") or "pending",
                "upload_status": _cell(row, col_index, "upload_status") or "pending",
                "acf_json": _cell(row, col_index, "acf_json"),
                "error": _cell(row, col_index, "error"),
                "scraped_at": _cell(row, col_index, "scraped_at"),
                "uploaded_at": _cell(row, col_index, "uploaded_at"),
            }
        )
    return jobs, wb, col_index


def save_job(wb, col_index: dict, job: dict) -> None:
    ws = wb.active
    row_idx = job["excel_row"]
    mapping = {
        "url": job.get("url", ""),
        "scrape_status": job.get("scrape_status", ""),
        "upload_status": job.get("upload_status", ""),
        "acf_json": job.get("acf_json", ""),
        "error": job.get("error", ""),
        "scraped_at": job.get("scraped_at", ""),
        "uploaded_at": job.get("uploaded_at", ""),
    }
    for field, value in mapping.items():
        if field in col_index:
            ws.cell(row=row_idx, column=col_index[field] + 1, value=value)


def persist(wb, path: Path, col_index: dict, job: dict) -> None:
    save_job(wb, col_index, job)
    wb.save(path)


def run_scrape(
    url: str,
    stem: str,
    outdir: Path,
    *,
    remove_logo: bool,
    rephrase_content: bool,
    selenium: bool,
    max_images: int,
    delay: float,
) -> Path:
    outdir.mkdir(parents=True, exist_ok=True)
    cmd = [
        sys.executable,
        str(SCRIPT_DIR / "scrapper.py"),
        "--url", url,
        "--output", stem,
        "--outdir", str(outdir),
        "--delay", str(delay),
        "--max-images", str(max_images),
    ]
    if remove_logo:
        cmd.append("--remove-logo")
    if rephrase_content:
        cmd.append("--rephrase-content")
    if selenium:
        cmd.append("--selenium")

    log.info("Scraping: %s", url)
    result = subprocess.run(cmd, cwd=str(SCRIPT_DIR), capture_output=True, text=True)
    if result.stdout:
        print(result.stdout, end="" if result.stdout.endswith("\n") else "\n")
    if result.returncode != 0:
        err = (result.stderr or result.stdout or "scrapper.py failed").strip()
        raise RuntimeError(err[-2000:])

    acf_path = outdir / f"{stem}_acf.json"
    if not acf_path.exists():
        raise RuntimeError(f"Expected ACF JSON not found: {acf_path}")
    return acf_path


def run_upload(
    acf_json: Path,
    *,
    wp_url: str,
    user: str,
    password: str,
    images_dir: str,
    dry_run: bool,
) -> None:
    cmd = [
        sys.executable,
        str(SCRIPT_DIR / "wp_uploader.py"),
        str(acf_json),
        "--wp-url", wp_url,
        "--user", user,
        "--password", password,
        "--max", "1",
    ]
    if images_dir:
        cmd.extend(["--images-dir", images_dir])
    if dry_run:
        cmd.append("--dry-run")

    log.info("Uploading: %s", acf_json.name)
    result = subprocess.run(cmd, cwd=str(SCRIPT_DIR), capture_output=True, text=True)
    if result.stdout:
        print(result.stdout, end="" if result.stdout.endswith("\n") else "\n")
    if result.returncode != 0:
        err = (result.stderr or result.stdout or "wp_uploader.py failed").strip()
        raise RuntimeError(err[-2000:])


def resolve_wp_credentials(args) -> tuple[str, str, str]:
    wp_url = args.wp_url or os.getenv("WP_URL", "")
    user = args.user or os.getenv("WP_USER", "")
    password = args.password or os.getenv("WP_PASSWORD", "")
    missing = [name for name, val in (("wp-url", wp_url), ("user", user), ("password", password)) if not val]
    if missing and not args.scrape_only:
        raise SystemExit(
            "WordPress credentials required for upload. Set WP_URL, WP_USER, WP_PASSWORD in .env "
            f"or pass {', '.join('--' + m for m in missing)}."
        )
    return wp_url, user, password


def should_scrape(job: dict, *, force: bool) -> bool:
    if force:
        return True
    return job["scrape_status"].lower() not in ("done", "skipped")


def should_upload(job: dict, *, force: bool) -> bool:
    if force:
        return True
    return job["upload_status"].lower() not in ("done", "skipped")


def process_jobs(args) -> int:
    excel_path = Path(args.excel).resolve()
    if not excel_path.exists():
        log.error("Spreadsheet not found: %s", excel_path)
        return 1

    jobs, wb, col_index = load_jobs(excel_path)
    outdir = Path(args.outdir).resolve()
    wp_url, user, password = resolve_wp_credentials(args)

    pending_scrape = sum(1 for j in jobs if should_scrape(j, force=args.force))
    pending_upload = sum(
        1 for j in jobs
        if should_upload(j, force=args.force) and (j["acf_json"] or should_scrape(j, force=False))
    )

    log.info("=" * 60)
    log.info("Spreadsheet  : %s", excel_path)
    log.info("Total rows   : %d", len(jobs))
    log.info("To scrape    : %d", pending_scrape if not args.upload_only else 0)
    log.info("To upload    : %d", pending_upload if not args.scrape_only else 0)
    log.info("Output dir   : %s", outdir)
    log.info("=" * 60)

    errors = 0
    for i, job in enumerate(jobs, 1):
        brand, model = job["brand"], job["model"]
        url = job["url"] or build_url(brand, model)
        stem = output_stem(brand, model)
        job["url"] = url

        log.info("[%d / %d]  %s %s", i, len(jobs), brand, model)

        # --- Scrape ---
        if not args.upload_only and should_scrape(job, force=args.force):
            job["error"] = ""
            try:
                acf_path = run_scrape(
                    url,
                    stem,
                    outdir,
                    remove_logo=args.remove_logo,
                    rephrase_content=args.rephrase_content,
                    selenium=args.selenium,
                    max_images=args.max_images,
                    delay=args.scrape_delay,
                )
                job["scrape_status"] = "done"
                job["acf_json"] = str(acf_path.resolve())
                job["scraped_at"] = now_str()
                log.info("  Scrape OK → %s", acf_path.name)
            except Exception as exc:
                job["scrape_status"] = "error"
                job["error"] = str(exc)[:500]
                log.error("  Scrape FAILED: %s", exc)
                errors += 1
                persist(wb, excel_path, col_index, job)
                continue
            persist(wb, excel_path, col_index, job)

        # --- Upload ---
        if not args.scrape_only and should_upload(job, force=args.force):
            acf_json = job.get("acf_json") or str((outdir / f"{stem}_acf.json").resolve())
            if not Path(acf_json).exists():
                job["upload_status"] = "error"
                job["error"] = f"ACF JSON not found: {acf_json}"
                log.error("  Upload skipped — no ACF JSON")
                errors += 1
                persist(wb, excel_path, col_index, job)
                continue

            job["error"] = ""
            try:
                run_upload(
                    Path(acf_json),
                    wp_url=wp_url,
                    user=user,
                    password=password,
                    images_dir=args.images_dir or str(outdir / "images"),
                    dry_run=args.dry_run,
                )
                job["upload_status"] = "done" if not args.dry_run else job["upload_status"]
                if not args.dry_run:
                    job["uploaded_at"] = now_str()
                log.info("  Upload OK")
            except Exception as exc:
                job["upload_status"] = "error"
                job["error"] = str(exc)[:500]
                log.error("  Upload FAILED: %s", exc)
                errors += 1
            persist(wb, excel_path, col_index, job)

    log.info("=" * 60)
    log.info("Finished. Errors: %d. Spreadsheet updated → %s", errors, excel_path)
    log.info("=" * 60)
    return 1 if errors else 0


def main():
    parser = argparse.ArgumentParser(
        description="Excel-driven car scrape + WordPress upload automation",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    parser.add_argument(
        "excel",
        nargs="?",
        help="Path to Excel file (.xlsx) with brand/model rows",
    )
    parser.add_argument(
        "--create-template",
        metavar="PATH",
        help="Create a starter Excel template and exit",
    )
    parser.add_argument("--outdir", default=str(DEFAULT_OUTDIR), help="Scrape output directory")
    parser.add_argument("--scrape-only", action="store_true", help="Only scrape, skip WordPress upload")
    parser.add_argument("--upload-only", action="store_true", help="Only upload rows that already have acf_json")
    parser.add_argument("--force", action="store_true", help="Re-process rows even if status is done")
    parser.add_argument("--dry-run", action="store_true", help="Dry-run WordPress upload (no writes)")

    parser.add_argument("--wp-url", default="", help="WordPress site URL (or WP_URL in .env)")
    parser.add_argument("--user", default="", help="WordPress username (or WP_USER in .env)")
    parser.add_argument("--password", default="", help="WordPress app password (or WP_PASSWORD in .env)")
    parser.add_argument("--images-dir", default="", help="Local images root for wp_uploader (default: <outdir>/images)")

    parser.add_argument("--remove-logo", action="store_true", help="Pass --remove-logo to scrapper.py")
    parser.add_argument("--rephrase-content", action="store_true", help="Pass --rephrase-content to scrapper.py")
    parser.add_argument("--selenium", action="store_true", help="Pass --selenium to scrapper.py")
    parser.add_argument("--max-images", type=int, default=5, help="Max gallery images per model (default: 5)")
    parser.add_argument("--scrape-delay", type=float, default=2.0, help="Delay between scrape requests")

    args = parser.parse_args()

    if args.create_template:
        create_template(Path(args.create_template))
        return 0

    if not args.excel:
        parser.error("excel file path is required (or use --create-template)")

    if args.scrape_only and args.upload_only:
        parser.error("--scrape-only and --upload-only cannot be used together")

    return process_jobs(args)


if __name__ == "__main__":
    sys.exit(main())