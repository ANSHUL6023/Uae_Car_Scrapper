#!/usr/bin/env python3
"""
automate_brands.py — Extract car URLs from brand pages or sitemaps to feed into automate.py

Reads an Excel sheet containing brand URLs (or sitemap URLs).
For each URL, it visits the page, extracts all individual car model URLs,
and generates a new Excel sheet formatted for `automate.py` (with brand, model, url, etc.).

Excel columns expected in input file (header row required):
  brand_url

Usage
-----
  # Create a starter spreadsheet for brands
  python automate_brands.py --create-template brands.xlsx

  # Generate cars.xlsx from brands.xlsx
  python automate_brands.py brands.xlsx --output cars.xlsx

  # Generate cars.xlsx and immediately run automate.py on it
  python automate_brands.py brands.xlsx --output cars.xlsx --run-automate --remove-logo --rephrase-content

"""

import argparse
import logging
import re
import sys
import subprocess
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    print("ERROR: openpyxl is required. Run: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-7s  %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger(__name__)

SCRIPT_DIR = Path(__file__).resolve().parent

# automate.py columns
CARS_COLUMNS = [
    "brand",
    "model",
    "url",
    "scrape_status",
    "upload_status",
    "acf_json",
    "error",
    "scraped_at",
    "uploaded_at",
]

def create_template(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "brands"
    ws.append(["brand_url"])
    ws.append(["https://www.zigwheels.ae/new-cars/hyundai"])
    ws.append(["https://www.zigwheels.ae/new-cars/toyota"])
    wb.save(path)
    log.info("Template created → %s", path.resolve())

def fetch(url: str) -> str:
    headers = {
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
        "Accept-Language": "en-US,en;q=0.5",
    }
    log.info("Fetching %s", url)
    resp = requests.get(url, headers=headers, timeout=20)
    resp.raise_for_status()
    return resp.text

def parse_sitemap_xml(xml_content: str) -> list[str]:
    soup = BeautifulSoup(xml_content, "xml")
    locs = soup.find_all("loc")
    urls = []
    for loc in locs:
        u = loc.text.strip()
        if re.search(r"/new-cars/[a-z0-9-]+/[a-z0-9-]+$", u, re.I):
            urls.append(u)
    return urls

def extract_cars_from_brand_page(html_content: str, base_url: str) -> list[str]:
    soup = BeautifulSoup(html_content, "html.parser")
    car_urls = set()
    for a in soup.find_all("a", href=True):
        href = a["href"]
        full_url = urljoin(base_url, href).split("#")[0].split("?")[0]
        
        # We are looking for URLs like https://www.zigwheels.ae/new-cars/hyundai/creta
        parsed = urlparse(full_url)
        path_parts = [p for p in parsed.path.strip("/").split("/") if p]
        
        if len(path_parts) == 3 and path_parts[0] == "new-cars":
            # Exclude non-car sub-pages that might slip through
            exclusions = {"dealers", "videos", "images", "news", "user-reviews", "faq"}
            if path_parts[2] not in exclusions:
                car_urls.add(full_url)
                
    return list(car_urls)

def extract_brand_model_from_url(url: str) -> tuple[str, str]:
    # https://www.zigwheels.ae/new-cars/brand/model
    path = urlparse(url).path.strip("/")
    parts = [p for p in path.split("/") if p]
    if len(parts) >= 3 and parts[0] == "new-cars":
        return parts[1], parts[2]
    return "unknown", "unknown"

def process_brands(args) -> int:
    input_path = Path(args.input_excel).resolve()
    if not input_path.exists():
        log.error("Input Excel not found: %s", input_path)
        return 1

    wb = load_workbook(input_path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        log.error("Spreadsheet is empty.")
        return 1

    headers = [str(h).strip().lower() if h else "" for h in rows[0]]
    if "brand_url" not in headers:
        log.error("Missing required column: brand_url")
        return 1

    col_idx = headers.index("brand_url")
    
    all_car_urls = set()
    
    for i, row in enumerate(rows[1:], start=2):
        brand_url = row[col_idx]
        if not brand_url:
            continue
        brand_url = str(brand_url).strip()
        
        try:
            content = fetch(brand_url)
            if brand_url.endswith(".xml"):
                car_urls = parse_sitemap_xml(content)
            else:
                car_urls = extract_cars_from_brand_page(content, brand_url)
                
            log.info("Found %d car URLs for %s", len(car_urls), brand_url)
            all_car_urls.update(car_urls)
            
        except Exception as e:
            log.error("Failed to process %s: %s", brand_url, e)

    if not all_car_urls:
        log.warning("No car URLs found across any brand.")
        return 0

    # Write to output cars.xlsx
    out_path = Path(args.output).resolve()
    out_wb = Workbook()
    out_ws = out_wb.active
    out_ws.title = "cars"
    out_ws.append(CARS_COLUMNS)
    
    for url in sorted(all_car_urls):
        brand, model = extract_brand_model_from_url(url)
        # Columns: brand, model, url, scrape_status, upload_status, acf_json, error, scraped_at, uploaded_at
        row = [
            brand,
            model,
            url,
            "pending",
            "pending",
            "",
            "",
            "",
            ""
        ]
        out_ws.append(row)
        
    out_wb.save(out_path)
    log.info("=" * 60)
    log.info("Successfully generated %s with %d car URLs.", out_path.name, len(all_car_urls))
    
    if args.run_automate:
        log.info("=" * 60)
        log.info("Running automate.py on %s", out_path.name)
        cmd = [sys.executable, str(SCRIPT_DIR / "automate.py"), str(out_path)]
        if args.remove_logo:
            cmd.append("--remove-logo")
        if args.rephrase_content:
            cmd.append("--rephrase-content")
        if args.scrape_only:
            cmd.append("--scrape-only")
        if args.upload_only:
            cmd.append("--upload-only")
            
        log.info("Command: %s", " ".join(cmd))
        subprocess.run(cmd, cwd=str(SCRIPT_DIR))
        
    return 0

def main():
    parser = argparse.ArgumentParser(
        description="Extract car URLs from brand URLs to generate an automate.py compatible Excel sheet",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=__doc__,
    )
    parser.add_argument("input_excel", nargs="?", help="Path to Excel file (.xlsx) with brand_url column")
    parser.add_argument("--create-template", metavar="PATH", help="Create a starter Excel template and exit")
    parser.add_argument("--output", default="cars.xlsx", help="Output Excel file (default: cars.xlsx)")
    
    # Arguments to pass-through to automate.py
    parser.add_argument("--run-automate", action="store_true", help="Automatically run automate.py on the generated output file")
    parser.add_argument("--remove-logo", action="store_true", help="Pass --remove-logo to automate.py")
    parser.add_argument("--rephrase-content", action="store_true", help="Pass --rephrase-content to automate.py")
    parser.add_argument("--scrape-only", action="store_true", help="Pass --scrape-only to automate.py")
    parser.add_argument("--upload-only", action="store_true", help="Pass --upload-only to automate.py")

    args = parser.parse_args()

    if args.create_template:
        create_template(Path(args.create_template))
        return 0

    if not args.input_excel:
        parser.error("input_excel file path is required (or use --create-template)")

    return process_brands(args)

if __name__ == "__main__":
    sys.exit(main())
